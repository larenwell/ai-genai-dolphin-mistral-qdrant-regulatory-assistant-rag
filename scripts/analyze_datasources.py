#!/usr/bin/env python3
"""
Análisis Completo de Fuentes de Datos
=====================================

Este script analiza todas las fuentes de datos en las carpetas:
- data/normativa/
- data/NFPA/
- data/normativa_extra/
- data/ingested/
- data/test/

Genera un reporte completo con:
- Cantidad de páginas por documento
- Detección de idioma (español/inglés)
- Tipo de archivo (PDF/DOCX)
- Ubicación (carpeta)
- Detección de duplicados
- Estadísticas generales

Autor: Laren Osorio Toribio
Fecha: 2025-01-05
"""

import os
import sys
import json
import hashlib
import time
from pathlib import Path
from typing import Dict, List, Tuple, Optional
from collections import defaultdict, Counter
import re

# Agregar el directorio src al path
project_root = Path(__file__).parent.parent
sys.path.append(str(project_root / "src"))

try:
    import PyPDF2
    import pymupdf as fitz
    from docx import Document
    import langdetect
    from langdetect import detect
except ImportError as e:
    print(f"❌ Error: Dependencias faltantes: {e}")
    print("💡 Instala las dependencias con: uv add PyPDF2 pymupdf python-docx langdetect")
    sys.exit(1)

class DataSourceAnalyzer:
    def __init__(self, data_root: str = None):
        """Inicializar el analizador de fuentes de datos."""
        if data_root is None:
            self.data_root = project_root / "data"
        else:
            self.data_root = Path(data_root)
        
        self.output_dir = project_root / "src" / "output" / "datasources"
        self.output_dir.mkdir(parents=True, exist_ok=True)
        
        # Configuración de carpetas a analizar
        self.folders_to_analyze = [
            "normativa",
            "NFPA", 
            "normativa_extra",
            "ingested",
            "test"
        ]
        
        # Resultados del análisis
        self.analysis_results = {
            "metadata": {
                "analysis_date": time.strftime("%Y-%m-%d %H:%M:%S"),
                "total_files": 0,
                "total_pages": 0,
                "folders_analyzed": self.folders_to_analyze
            },
            "files": [],
            "duplicates": [],
            "statistics": {},
            "language_analysis": {},
            "file_type_analysis": {},
            "folder_analysis": {}
        }
        
        # Cache para detección de duplicados
        self.file_hashes = {}
        self.duplicate_groups = defaultdict(list)
    
    def print_stage_title(self, title: str, stage_number: int = None):
        """Imprimir título de etapa con separadores visuales."""
        if stage_number:
            print(f"\n{'='*80}")
            print(f"🚀 ETAPA {stage_number}: {title}")
            print(f"{'='*80}")
        else:
            print(f"\n{'='*80}")
            print(f"🎯 {title}")
            print(f"{'='*80}")
    
    def print_sub_stage(self, title: str):
        """Imprimir subtítulo de etapa."""
        print(f"\n📋 {title}")
        print(f"{'-'*60}")
    
    def calculate_file_hash(self, file_path: Path) -> str:
        """Calcular hash MD5 del archivo para detección de duplicados."""
        try:
            hash_md5 = hashlib.md5()
            with open(file_path, "rb") as f:
                for chunk in iter(lambda: f.read(4096), b""):
                    hash_md5.update(chunk)
            return hash_md5.hexdigest()
        except Exception as e:
            print(f"   ⚠️ Error calculando hash para {file_path}: {e}")
            return ""
    
    def detect_language_from_filename(self, filename: str) -> str:
        """Detectar idioma basado en el nombre del archivo."""
        filename_lower = filename.lower()
        
        # Indicadores de español
        spanish_indicators = [
            'norma', 'decreto', 'resolucion', 'reglamento', 'ley', 'ds', 'rm', 'rcd',
            'nro', 'n°', 'edicion', 'ed', 'anexos', 'suministro', 'utilizacion',
            'comercializacion', 'modifica', 'modificacion', 'seguridad', 'transporte',
            'instalaciones', 'distribucion', 'venta', 'mineria', 'construccion',
            'vivienda', 'edificios', 'sistemas', 'gestion', 'procesos'
        ]
        
        # Indicadores de inglés
        english_indicators = [
            'nfpa', 'code', 'standard', 'handbook', 'guide', 'manual', 'specification',
            'requirements', 'installation', 'maintenance', 'inspection', 'testing',
            'fire', 'safety', 'protection', 'system', 'equipment', 'materials',
            'construction', 'building', 'electrical', 'mechanical', 'chemical'
        ]
        
        spanish_count = sum(1 for indicator in spanish_indicators if indicator in filename_lower)
        english_count = sum(1 for indicator in english_indicators if indicator in filename_lower)
        
        if spanish_count > english_count:
            return "español"
        elif english_count > spanish_count:
            return "inglés"
        else:
            return "indeterminado"
    
    def detect_language_from_content(self, file_path: Path, max_pages: int = 3) -> str:
        """Detectar idioma del contenido del archivo."""
        try:
            # Leer solo las primeras páginas para análisis de idioma
            text_sample = ""
            
            if file_path.suffix.lower() == '.pdf':
                doc = fitz.open(file_path)
                for page_num in range(min(max_pages, len(doc))):
                    page = doc[page_num]
                    text_sample += page.get_text()
                doc.close()
            elif file_path.suffix.lower() == '.docx':
                doc = Document(file_path)
                for paragraph in doc.paragraphs[:10]:  # Primeros 10 párrafos
                    text_sample += paragraph.text + " "
            
            if len(text_sample.strip()) < 50:
                return "insuficiente_texto"
            
            # Detectar idioma usando langdetect
            try:
                detected_lang = detect(text_sample)
                lang_map = {
                    'es': 'español',
                    'en': 'inglés',
                    'pt': 'portugués'
                }
                return lang_map.get(detected_lang, detected_lang)
            except:
                return "error_deteccion"
                
        except Exception as e:
            print(f"   ⚠️ Error detectando idioma en {file_path}: {e}")
            return "error"
    
    def get_page_count(self, file_path: Path) -> int:
        """Obtener número de páginas del archivo."""
        try:
            if file_path.suffix.lower() == '.pdf':
                # Usar PyMuPDF para mayor precisión
                doc = fitz.open(file_path)
                page_count = len(doc)
                doc.close()
                return page_count
            elif file_path.suffix.lower() == '.docx':
                # Para DOCX, estimar páginas basado en contenido
                doc = Document(file_path)
                paragraph_count = len(doc.paragraphs)
                # Estimación aproximada: 1 página por cada 20-30 párrafos
                estimated_pages = max(1, paragraph_count // 25)
                return estimated_pages
            else:
                return 1
        except Exception as e:
            print(f"   ⚠️ Error contando páginas en {file_path}: {e}")
            return 0
    
    def get_file_size_mb(self, file_path: Path) -> float:
        """Obtener tamaño del archivo en MB."""
        try:
            size_bytes = file_path.stat().st_size
            return round(size_bytes / (1024 * 1024), 2)
        except Exception as e:
            print(f"   ⚠️ Error obteniendo tamaño de {file_path}: {e}")
            return 0.0
    
    def analyze_file(self, file_path: Path, folder_name: str) -> Dict:
        """Analizar un archivo individual."""
        print(f"   📄 Analizando: {file_path.name}")
        
        # Información básica
        file_info = {
            "filename": file_path.name,
            "folder": folder_name,
            "full_path": str(file_path),
            "file_extension": file_path.suffix.lower(),
            "file_size_mb": self.get_file_size_mb(file_path),
            "page_count": self.get_page_count(file_path),
            "file_hash": self.calculate_file_hash(file_path),
            "analysis_timestamp": time.strftime("%Y-%m-%d %H:%M:%S")
        }
        
        # Determinar idioma final
        language_from_content = self.detect_language_from_content(file_path)
        language_from_filename = self.detect_language_from_filename(file_path.name)
        
        if language_from_content not in ["error", "insuficiente_texto", "error_deteccion"]:
            file_info["final_language"] = language_from_content
        else:
            file_info["final_language"] = language_from_filename
        
        # Determinar tipo de archivo
        if file_info["file_extension"] == ".pdf":
            file_info["file_type"] = "PDF"
        elif file_info["file_extension"] == ".docx":
            file_info["file_type"] = "DOCX"
        else:
            file_info["file_type"] = "OTRO"
        
        return file_info
    
    def detect_duplicates(self):
        """Detectar archivos duplicados basado en hash MD5."""
        self.print_sub_stage("DETECCIÓN DE DUPLICADOS")
        
        # Agrupar archivos por hash
        for file_info in self.analysis_results["files"]:
            file_hash = file_info["file_hash"]
            if file_hash:
                self.duplicate_groups[file_hash].append(file_info)
        
        # Identificar grupos de duplicados
        duplicate_groups = {hash_val: files for hash_val, files in self.duplicate_groups.items() 
                          if len(files) > 1}
        
        self.analysis_results["duplicates"] = []
        for hash_val, files in duplicate_groups.items():
            duplicate_group = {
                "hash": hash_val,
                "count": len(files),
                "files": files,
                "total_size_mb": sum(f["file_size_mb"] for f in files),
                "wasted_space_mb": (len(files) - 1) * files[0]["file_size_mb"]
            }
            self.analysis_results["duplicates"].append(duplicate_group)
        
        print(f"   🔍 Grupos de duplicados encontrados: {len(duplicate_groups)}")
        total_duplicates = sum(len(files) - 1 for files in duplicate_groups.values())
        print(f"   📊 Archivos duplicados: {total_duplicates}")
        
        if duplicate_groups:
            wasted_space = sum(group["wasted_space_mb"] for group in self.analysis_results["duplicates"])
            print(f"   💾 Espacio desperdiciado: {wasted_space:.2f} MB")
    
    def generate_statistics(self):
        """Generar estadísticas del análisis."""
        self.print_sub_stage("GENERACIÓN DE ESTADÍSTICAS")
        
        files = self.analysis_results["files"]
        
        # Estadísticas generales
        self.analysis_results["statistics"] = {
            "total_files": len(files),
            "total_pages": sum(f["page_count"] for f in files),
            "total_size_mb": sum(f["file_size_mb"] for f in files),
            "average_file_size_mb": round(sum(f["file_size_mb"] for f in files) / len(files), 2) if files else 0,
            "average_pages_per_file": round(sum(f["page_count"] for f in files) / len(files), 1) if files else 0
        }
        
        # Análisis por idioma
        language_counts = Counter(f["final_language"] for f in files)
        self.analysis_results["language_analysis"] = {
            "distribution": dict(language_counts),
            "spanish_files": language_counts.get("español", 0),
            "english_files": language_counts.get("inglés", 0),
            "other_files": sum(count for lang, count in language_counts.items() 
                             if lang not in ["español", "inglés"])
        }
        
        # Análisis por tipo de archivo
        file_type_counts = Counter(f["file_type"] for f in files)
        self.analysis_results["file_type_analysis"] = {
            "distribution": dict(file_type_counts),
            "pdf_files": file_type_counts.get("PDF", 0),
            "docx_files": file_type_counts.get("DOCX", 0),
            "other_files": file_type_counts.get("OTRO", 0)
        }
        
        # Análisis por carpeta
        folder_stats = defaultdict(lambda: {"count": 0, "pages": 0, "size_mb": 0})
        for f in files:
            folder = f["folder"]
            folder_stats[folder]["count"] += 1
            folder_stats[folder]["pages"] += f["page_count"]
            folder_stats[folder]["size_mb"] += f["file_size_mb"]
        
        self.analysis_results["folder_analysis"] = dict(folder_stats)
        
        print(f"   📊 Archivos analizados: {self.analysis_results['statistics']['total_files']}")
        print(f"   📄 Total de páginas: {self.analysis_results['statistics']['total_pages']}")
        print(f"   💾 Tamaño total: {self.analysis_results['statistics']['total_size_mb']:.2f} MB")
        print(f"   🌐 Archivos en español: {self.analysis_results['language_analysis']['spanish_files']}")
        print(f"   🌐 Archivos en inglés: {self.analysis_results['language_analysis']['english_files']}")
        print(f"   📄 Archivos PDF: {self.analysis_results['file_type_analysis']['pdf_files']}")
        print(f"   📝 Archivos DOCX: {self.analysis_results['file_type_analysis']['docx_files']}")
    
    def analyze_folder(self, folder_name: str) -> List[Dict]:
        """Analizar una carpeta específica."""
        folder_path = self.data_root / folder_name
        
        if not folder_path.exists():
            print(f"   ⚠️ Carpeta {folder_name} no encontrada")
            return []
        
        print(f"   📁 Analizando carpeta: {folder_name}")
        
        # Obtener archivos soportados (incluyendo subcarpetas)
        # Buscar tanto en mayúsculas como en minúsculas
        supported_extensions = ['.pdf', '.PDF', '.docx', '.DOCX']
        files = []
        
        for ext in supported_extensions:
            # Buscar archivos recursivamente en subcarpetas
            files.extend(folder_path.glob(f"**/*{ext}"))
        
        # Eliminar duplicados manteniendo el orden
        files = list(dict.fromkeys(files))
        
        print(f"   📊 Archivos encontrados: {len(files)}")
        
        folder_files = []
        for file_path in files:
            try:
                file_info = self.analyze_file(file_path, folder_name)
                folder_files.append(file_info)
            except Exception as e:
                print(f"   ❌ Error analizando {file_path}: {e}")
                continue
        
        return folder_files
    
    def save_results(self):
        """Guardar resultados del análisis."""
        self.print_sub_stage("GUARDANDO RESULTADOS")
        
        timestamp = time.strftime("%Y%m%d_%H%M%S")
        
        # Guardar JSON completo
        json_file = self.output_dir / f"datasources_analysis_{timestamp}.json"
        with open(json_file, "w", encoding="utf-8") as f:
            json.dump(self.analysis_results, f, indent=2, ensure_ascii=False)
        
        # Guardar reporte en texto
        txt_file = self.output_dir / f"datasources_report_{timestamp}.txt"
        self.generate_text_report(txt_file)
        
        # Guardar CSV de archivos
        csv_file = self.output_dir / f"datasources_files_{timestamp}.csv"
        self.generate_csv_report(csv_file)
        
        print(f"   📁 Resultados guardados en: {self.output_dir}")
        print(f"   📄 JSON: {json_file.name}")
        print(f"   📄 TXT: {txt_file.name}")
        print(f"   📄 CSV: {csv_file.name}")
    
    def generate_text_report(self, output_file: Path):
        """Generar reporte en formato texto."""
        with open(output_file, "w", encoding="utf-8") as f:
            f.write("ANÁLISIS COMPLETO DE FUENTES DE DATOS\n")
            f.write("=" * 50 + "\n\n")
            
            # Metadatos
            f.write(f"Fecha de análisis: {self.analysis_results['metadata']['analysis_date']}\n")
            f.write(f"Carpetas analizadas: {', '.join(self.analysis_results['metadata']['folders_analyzed'])}\n")
            f.write(f"Total de archivos: {self.analysis_results['statistics']['total_files']}\n")
            f.write(f"Total de páginas: {self.analysis_results['statistics']['total_pages']}\n")
            f.write(f"Tamaño total: {self.analysis_results['statistics']['total_size_mb']:.2f} MB\n\n")
            
            # Análisis por idioma
            f.write("DISTRIBUCIÓN POR IDIOMA\n")
            f.write("-" * 30 + "\n")
            for lang, count in self.analysis_results['language_analysis']['distribution'].items():
                f.write(f"{lang}: {count} archivos\n")
            f.write("\n")
            
            # Análisis por tipo de archivo
            f.write("DISTRIBUCIÓN POR TIPO DE ARCHIVO\n")
            f.write("-" * 35 + "\n")
            for file_type, count in self.analysis_results['file_type_analysis']['distribution'].items():
                f.write(f"{file_type}: {count} archivos\n")
            f.write("\n")
            
            # Análisis por carpeta
            f.write("ANÁLISIS POR CARPETA\n")
            f.write("-" * 20 + "\n")
            for folder, stats in self.analysis_results['folder_analysis'].items():
                f.write(f"{folder}:\n")
                f.write(f"  Archivos: {stats['count']}\n")
                f.write(f"  Páginas: {stats['pages']}\n")
                f.write(f"  Tamaño: {stats['size_mb']:.2f} MB\n\n")
            
            # Duplicados
            if self.analysis_results['duplicates']:
                f.write("ARCHIVOS DUPLICADOS\n")
                f.write("-" * 20 + "\n")
                for i, group in enumerate(self.analysis_results['duplicates'], 1):
                    f.write(f"Grupo {i} ({group['count']} archivos):\n")
                    for file_info in group['files']:
                        f.write(f"  - {file_info['filename']} ({file_info['folder']})\n")
                    f.write(f"  Espacio desperdiciado: {group['wasted_space_mb']:.2f} MB\n\n")
            else:
                f.write("No se encontraron archivos duplicados.\n\n")
    
    def generate_csv_report(self, output_file: Path):
        """Generar reporte en formato CSV."""
        import csv
        
        with open(output_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            
            # Encabezados
            headers = [
                "filename", "folder", "file_type", "file_size_mb", "page_count",
                "final_language",
                "file_hash", "full_path"
            ]
            writer.writerow(headers)
            
            # Datos
            for file_info in self.analysis_results["files"]:
                row = [
                    file_info["filename"],
                    file_info["folder"],
                    file_info["file_type"],
                    file_info["file_size_mb"],
                    file_info["page_count"],
                    file_info["final_language"],
                    file_info["file_hash"],
                    file_info["full_path"]
                ]
                writer.writerow(row)
    
    def run_analysis(self):
        """Ejecutar análisis completo."""
        self.print_stage_title("ANÁLISIS DE FUENTES DE DATOS", 1)
        
        print(f"📁 Directorio de datos: {self.data_root}")
        print(f"📁 Directorio de salida: {self.output_dir}")
        print(f"📂 Carpetas a analizar: {', '.join(self.folders_to_analyze)}")
        
        # Analizar cada carpeta
        all_files = []
        for folder_name in self.folders_to_analyze:
            self.print_sub_stage(f"ANALIZANDO CARPETA: {folder_name}")
            folder_files = self.analyze_folder(folder_name)
            all_files.extend(folder_files)
            print(f"   ✅ Archivos procesados: {len(folder_files)}")
        
        self.analysis_results["files"] = all_files
        self.analysis_results["metadata"]["total_files"] = len(all_files)
        
        # Detectar duplicados
        self.detect_duplicates()
        
        # Generar estadísticas
        self.generate_statistics()
        
        # Guardar resultados
        self.save_results()
        
        self.print_stage_title("ANÁLISIS COMPLETADO")
        print(f"✅ Análisis completado exitosamente")
        print(f"📊 Total de archivos analizados: {len(all_files)}")
        print(f"🔍 Duplicados encontrados: {len(self.analysis_results['duplicates'])}")
        print(f"📁 Resultados guardados en: {self.output_dir}")

def main():
    """Función principal."""
    print("🚀 Iniciando Análisis de Fuentes de Datos")
    print("=" * 50)
    
    try:
        analyzer = DataSourceAnalyzer()
        analyzer.run_analysis()
    except KeyboardInterrupt:
        print("\n⚠️ Análisis interrumpido por el usuario")
        sys.exit(1)
    except Exception as e:
        print(f"\n❌ Error durante el análisis: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)

if __name__ == "__main__":
    main()
